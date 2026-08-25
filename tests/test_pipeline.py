"""Tests du pipeline bout-en-bout (CDC 5.1 / 5.8)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from generateur.synthetique import generer_corpus
from persistance.base import DocumentDB, DrapeauRelectureDB, FactureDB, JournalAuditDB
from persistance.pipeline import exporter_csv, exporter_excel, traiter_dossier


@pytest.fixture(scope="module")
def corpus(tmp_path_factory) -> Path:
    dossier = tmp_path_factory.mktemp("corpus_pipeline")
    generer_corpus(dossier, nb_coherentes=8, nb_fautives=4, graine=33)
    return dossier


def test_lot_complet_statuts_et_drapeaux(corpus, tmp_path):
    url = f"sqlite:///{tmp_path / 'test.db'}"
    bilan = traiter_dossier(corpus, url)

    assert bilan["auto_accepte"] + bilan["a_relire"] == 12
    assert bilan["a_relire"] >= 4  # les 4 fautes injectees, au minimum
    assert bilan["echec"] == 0 and bilan["doublon"] == 0

    engine = create_engine(url)
    with Session(engine) as session:
        assert session.scalar(select(func.count()).select_from(DocumentDB)) == 12
        # Chaque document a_relire porte au moins un drapeau explicite (EF-24)
        for doc in session.scalars(select(DocumentDB).where(DocumentDB.statut == "a_relire")):
            assert doc.facture is not None and len(doc.facture.drapeaux) >= 1
        # L'extraction brute est conservee (EF-26)
        un_doc = session.scalar(select(DocumentDB).where(DocumentDB.statut == "auto_accepte"))
        assert un_doc is not None and un_doc.extraction_brute
        # Le journal d'audit a trace le lot (ENF-07)
        assert session.scalar(select(func.count()).select_from(JournalAuditDB)) >= 12


def test_idempotence_rejouer_un_lot(corpus, tmp_path):
    """EF-28 : rejouer le meme lot ne cree aucun doublon."""
    url = f"sqlite:///{tmp_path / 'test.db'}"
    traiter_dossier(corpus, url)
    bilan2 = traiter_dossier(corpus, url)

    assert bilan2["doublon"] == 12
    assert bilan2["auto_accepte"] == 0 and bilan2["a_relire"] == 0

    engine = create_engine(url)
    with Session(engine) as session:
        assert session.scalar(select(func.count()).select_from(DocumentDB)) == 12
        assert session.scalar(select(func.count()).select_from(FactureDB)) == 12


def test_fichier_corrompu_ne_bloque_pas_le_lot(corpus, tmp_path):
    """EF-04 / ENF-06 : un PDF invalide est marque echec, le lot continue."""
    entree = tmp_path / "entree"
    entree.mkdir()
    for pdf in sorted(corpus.rglob("*.pdf"))[:3]:
        (entree / pdf.name).write_bytes(pdf.read_bytes())
    (entree / "corrompu.pdf").write_bytes(b"ceci n'est pas un pdf")

    bilan = traiter_dossier(entree, f"sqlite:///{tmp_path / 'test.db'}")
    assert bilan["echec"] == 1
    assert bilan["auto_accepte"] + bilan["a_relire"] == 3


def test_exports_csv_et_excel(corpus, tmp_path):
    """EF-29 : CSV UTF-8 BOM point-virgule + Excel deux feuilles."""
    url = f"sqlite:///{tmp_path / 'test.db'}"
    traiter_dossier(corpus, url)

    c_entetes, c_lignes = exporter_csv(url, tmp_path / "exports")
    brut = c_entetes.read_bytes()
    assert brut.startswith(b"\xef\xbb\xbf")            # BOM UTF-8
    premiere = brut.decode("utf-8-sig").splitlines()[0]
    assert ";" in premiere and "facture_id" in premiere
    assert len(c_entetes.read_text(encoding="utf-8-sig").splitlines()) == 13  # 12 + en-tete
    assert len(c_lignes.read_text(encoding="utf-8-sig").splitlines()) > 1

    xl = exporter_excel(url, tmp_path / "exports" / "factures.xlsx")
    classeur = load_workbook(xl)
    assert classeur.sheetnames == ["entetes", "lignes"]
    assert classeur["entetes"].max_row == 13
